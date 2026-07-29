"""
Script para gerar o questionário de levantamento de requisitos do projeto FacePhoto.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_background(cell, hex_color: str):
    """Define a cor de fundo de uma célula da tabela."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_heading(doc: Document, text: str, level: int = 1):
    """Adiciona um título ao documento com estilo personalizado."""
    heading = doc.add_heading(text, level=level)
    run = heading.runs[0]

    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x2E, 0x6D, 0xB0)


def add_question_block(
    doc: Document,
    question_number: str,
    question_text: str,
    options: list[str] | None = None,
    answer_lines: int = 3,
):
    """
    Adiciona um bloco de pergunta com campo de resposta ao documento.

    Args:
        doc: Documento Word sendo construído.
        question_number: Número/identificador da pergunta (ex: '1.1').
        question_text: Texto da pergunta.
        options: Lista de opções (quando for múltipla escolha). None para discursiva.
        answer_lines: Número de linhas em branco para resposta discursiva.
    """
    # Número + texto da pergunta
    para = doc.add_paragraph()
    run_num = para.add_run(f"{question_number}. ")
    run_num.bold = True
    run_num.font.size = Pt(11)
    run_num.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

    run_q = para.add_run(question_text)
    run_q.bold = True
    run_q.font.size = Pt(11)
    run_q.font.color.rgb = RGBColor(0x1E, 0x1E, 0x1E)

    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(4)

    if options:
        # Opções de múltipla escolha com checkbox visual
        for option in options:
            opt_para = doc.add_paragraph(style="List Bullet")
            opt_para.paragraph_format.left_indent = Inches(0.3)
            opt_para.paragraph_format.space_before = Pt(2)
            opt_para.paragraph_format.space_after = Pt(2)
            run_opt = opt_para.add_run(f"☐  {option}")
            run_opt.font.size = Pt(10.5)
        # Linha extra para observações
        obs_para = doc.add_paragraph()
        obs_run = obs_para.add_run("Observações / outros: ")
        obs_run.font.size = Pt(10)
        obs_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        obs_para.add_run("_" * 60)
        obs_para.paragraph_format.space_before = Pt(4)
        obs_para.paragraph_format.space_after = Pt(10)
    else:
        # Campo discursivo com linhas
        for _ in range(answer_lines):
            line_para = doc.add_paragraph()
            line_run = line_para.add_run("_" * 90)
            line_run.font.size = Pt(10)
            line_run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            line_para.paragraph_format.space_before = Pt(1)
            line_para.paragraph_format.space_after = Pt(1)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_section_divider(doc: Document):
    """Adiciona uma linha separadora horizontal."""
    para = doc.add_paragraph()
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E6DB0")
    pBdr.append(bottom)
    pPr.append(pBdr)
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(10)


def build_document(output_path: str):
    """Constrói e salva o documento de levantamento de requisitos."""
    doc = Document()

    # Configuração de margens
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # Fonte padrão
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ── CABEÇALHO ─────────────────────────────────────────────────────────────
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_para.add_run("📋  Questionário de Levantamento de Requisitos")
    run_title.bold = True
    run_title.font.size = Pt(18)
    run_title.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = subtitle.add_run("Projeto: FacePhoto — Reconhecimento Facial em Fotos")
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    run_sub.italic = True

    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_date = date_para.add_run("Data de preenchimento: _____ / _____ / _______")
    run_date.font.size = Pt(10)
    run_date.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph()

    # Instrução inicial
    inst_para = doc.add_paragraph()
    run_inst = inst_para.add_run(
        "ℹ️  Instruções: "
    )
    run_inst.bold = True
    run_inst.font.size = Pt(10.5)
    run_inst.font.color.rgb = RGBColor(0x2E, 0x6D, 0xB0)
    run_inst2 = inst_para.add_run(
        "Responda cada pergunta no campo indicado. "
        "Para perguntas de múltipla escolha, marque (✔) as opções que se aplicam. "
        "Não há resposta certa ou errada — o objetivo é entender o seu cenário real."
    )
    run_inst2.font.size = Pt(10.5)
    run_inst2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    inst_para.paragraph_format.space_after = Pt(14)

    add_section_divider(doc)

    # ── BLOCO 1: CONTEXTO E PROPÓSITO ─────────────────────────────────────────
    add_heading(doc, "1. Contexto e Propósito", level=1)

    add_question_block(
        doc,
        "1.1",
        "Para que você vai usar esse programa? Descreva o caso de uso principal.",
        answer_lines=3,
    )
    add_question_block(
        doc,
        "1.2",
        "Qual é o objetivo final ao usar o programa?",
        options=[
            "Organizar álbuns de fotos pessoais ou familiares",
            "Encontrar fotos de um evento específico (casamento, festa, etc.)",
            "Uso profissional (fotógrafo, empresa, segurança)",
            "Pesquisa / aprendizado pessoal",
            "Outro (descreva nas observações)",
        ],
    )
    add_question_block(
        doc,
        "1.3",
        "Quem vai usar o programa?",
        options=[
            "Somente eu (usuário técnico)",
            "Somente eu (usuário leigo — precisa ser simples)",
            "Um pequeno grupo de pessoas técnicas",
            "Outras pessoas leigas também vão usar",
        ],
    )

    add_section_divider(doc)

    # ── BLOCO 2: FOTOS DE REFERÊNCIA ──────────────────────────────────────────
    add_heading(doc, "2. Fotos de Referência (a pessoa que será buscada)", level=1)

    add_question_block(
        doc,
        "2.1",
        "Quantas fotos de referência você normalmente terá de uma mesma pessoa?",
        options=[
            "1 foto apenas",
            "2 a 5 fotos",
            "6 a 20 fotos",
            "Mais de 20 fotos",
        ],
    )
    add_question_block(
        doc,
        "2.2",
        "Como são as fotos de referência? Marque tudo que se aplica.",
        options=[
            "Fotos frontais nítidas (selfies, documentos, retratos)",
            "Fotos em grupo (a pessoa aparece entre outras)",
            "Fotos com variação de ângulo, luz ou expressão",
            "Fotos de baixa qualidade ou antigas",
        ],
    )
    add_question_block(
        doc,
        "2.3",
        "Você vai buscar por uma pessoa por vez ou por várias pessoas ao mesmo tempo?",
        options=[
            "Sempre uma pessoa por vez",
            "Às vezes precisarei buscar por 2 ou 3 pessoas ao mesmo tempo",
            "Quero poder buscar por muitas pessoas simultaneamente",
        ],
    )

    add_section_divider(doc)

    # ── BLOCO 3: FOTOS A SEREM ANALISADAS ────────────────────────────────────
    add_heading(doc, "3. Fotos a Serem Analisadas (o acervo de busca)", level=1)

    add_question_block(
        doc,
        "3.1",
        "De onde virão as fotos que o programa vai analisar? Marque tudo que se aplica.",
        options=[
            "Uma pasta local no computador (Windows Explorer)",
            "Um HD externo ou pendrive",
            "Uma pasta de rede (NAS, servidor local)",
            "Fotos importadas da internet ou redes sociais",
            "Outros (descreva nas observações)",
        ],
    )
    add_question_block(
        doc,
        "3.2",
        "Qual o volume estimado de fotos a serem analisadas por execução?",
        options=[
            "Até 100 fotos",
            "100 a 500 fotos",
            "500 a 5.000 fotos",
            "Mais de 5.000 fotos",
        ],
    )
    add_question_block(
        doc,
        "3.3",
        "As fotos estão organizadas em subpastas ou tudo em uma pasta única?",
        options=[
            "Tudo em uma única pasta",
            "Organizadas em subpastas (ex: por ano, evento, data)",
            "Misturado — sem organização definida",
        ],
    )
    add_question_block(
        doc,
        "3.4",
        "Quais formatos de imagem você costuma ter? Marque os que se aplicam.",
        options=[
            "JPG / JPEG",
            "PNG",
            "HEIC (iPhone)",
            "RAW (câmera profissional)",
            "BMP, TIFF ou outros",
        ],
    )

    add_section_divider(doc)

    # ── BLOCO 4: INTERFACE DO PROGRAMA ────────────────────────────────────────
    add_heading(doc, "4. Interface e Experiência do Usuário", level=1)

    add_question_block(
        doc,
        "4.1",
        "Como você prefere interagir com o programa?",
        options=[
            "Linha de comando (terminal) — estou confortável com isso",
            "Interface gráfica (janelas, botões, arrastar e soltar)",
            "Interface web (acessar pelo navegador, como um site local)",
            "Não tenho preferência",
        ],
    )
    add_question_block(
        doc,
        "4.2",
        "Em qual sistema operacional o programa será usado?",
        options=[
            "Windows (somente)",
            "macOS (somente)",
            "Linux (somente)",
            "Precisa funcionar em mais de um sistema",
        ],
    )
    add_question_block(
        doc,
        "4.3",
        "Você gostaria de acompanhar o progresso do processamento em tempo real "
        "(ex: barra de progresso, log de fotos analisadas)?",
        options=[
            "Sim, é importante ver o progresso",
            "Não precisa, pode rodar em silêncio e me avisar quando terminar",
            "Indiferente",
        ],
    )

    add_section_divider(doc)

    # ── BLOCO 5: SAÍDA E RESULTADOS ───────────────────────────────────────────
    add_heading(doc, "5. Saída e Resultados", level=1)

    add_question_block(
        doc,
        "5.1",
        "O que você quer que o programa faça com as fotos encontradas? Marque tudo que se aplica.",
        options=[
            "Copiar as fotos para uma pasta de destino (arquivo original preservado)",
            "Mover as fotos (remover da origem)",
            "Apenas listar quais fotos foram encontradas (sem copiar nada)",
            "Gerar um relatório (PDF ou HTML) com as fotos encontradas",
            "Criar subpastas separadas por pessoa encontrada",
        ],
    )
    add_question_block(
        doc,
        "5.2",
        "Quando uma foto contém a pessoa buscada junto com outras pessoas, o que deve acontecer?",
        options=[
            "Incluir a foto normalmente (a pessoa buscada está lá, não importa quem mais)",
            "Separar em uma pasta diferente (ex: 'grupo') para revisão",
            "Não incluir — só quero fotos onde a pessoa apareça sozinha ou em destaque",
        ],
    )
    add_question_block(
        doc,
        "5.3",
        "O que fazer quando o programa não tem certeza se é a pessoa certa?",
        options=[
            "Incluir assim mesmo (prefiro não perder nenhuma foto)",
            "Ignorar (prefiro ter só resultados com alta confiança)",
            "Copiar para uma pasta separada de 'revisão manual'",
        ],
    )
    add_question_block(
        doc,
        "5.4",
        "Você quer que os arquivos copiados mantenham os metadados originais "
        "(data de criação, localização GPS, câmera usada)?",
        options=[
            "Sim, é importante preservar os metadados",
            "Não me importo",
            "Não sei o que são metadados",
        ],
    )

    add_section_divider(doc)

    # ── BLOCO 6: DESEMPENHO E AMBIENTE ────────────────────────────────────────
    add_heading(doc, "6. Desempenho e Ambiente Técnico", level=1)

    add_question_block(
        doc,
        "6.1",
        "Seu computador tem placa de vídeo (GPU) dedicada? "
        "(Ex: NVIDIA GeForce, AMD Radeon — não integrada da Intel)",
        options=[
            "Sim, tenho GPU dedicada NVIDIA",
            "Sim, tenho GPU dedicada AMD",
            "Não sei / Tenho apenas GPU integrada",
            "Não tenho GPU dedicada",
        ],
    )
    add_question_block(
        doc,
        "6.2",
        "Qual é a sua expectativa de tempo de processamento?",
        options=[
            "Precisa ser rápido (minutos para centenas de fotos)",
            "Aceito que demore (pode levar horas para grandes volumes)",
            "Não tenho uma expectativa definida",
        ],
    )
    add_question_block(
        doc,
        "6.3",
        "Você prefere que o programa funcione completamente offline (sem internet)?",
        options=[
            "Sim, 100% offline — as fotos são privadas e não devem sair do meu computador",
            "Aceito usar serviços de nuvem se isso melhorar a precisão",
            "Não tenho preferência",
        ],
    )

    add_section_divider(doc)

    # ── BLOCO 7: PRECISÃO E COMPORTAMENTO ────────────────────────────────────
    add_heading(doc, "7. Precisão e Comportamento do Reconhecimento", level=1)

    add_question_block(
        doc,
        "7.1",
        "O que é mais crítico para você?",
        options=[
            "Não perder nenhuma foto (mesmo que inclua alguns falsos positivos — "
            "fotos de pessoas parecidas)",
            "Ter certeza que tudo encontrado é realmente a pessoa certa "
            "(mesmo que perca algumas fotos)",
            "Equilíbrio entre os dois",
        ],
    )
    add_question_block(
        doc,
        "7.2",
        "O programa precisa funcionar bem com fotos antigas ou de baixa qualidade "
        "(fotos digitalizadas, câmeras antigas, pouca iluminação)?",
        options=[
            "Sim, tenho muitas fotos antigas que precisam ser analisadas",
            "Não, minhas fotos são recentes e de boa qualidade",
            "Tenho uma mistura dos dois",
        ],
    )

    add_section_divider(doc)

    # ── BLOCO 8: PRIVACIDADE ─────────────────────────────────────────────────
    add_heading(doc, "8. Privacidade e Uso Ético", level=1)

    add_question_block(
        doc,
        "8.1",
        "As fotos que serão analisadas contêm conteúdo sensível ou privado?",
        options=[
            "Sim — são fotos pessoais/familiares confidenciais",
            "Parcialmente — algumas sim, algumas não",
            "Não — são fotos de eventos abertos ou de uso geral",
        ],
    )
    add_question_block(
        doc,
        "8.2",
        "O uso pretendido do programa é de natureza pessoal/privada?",
        options=[
            "Sim, uso estritamente pessoal (organizar minhas próprias fotos)",
            "Uso profissional (fotógrafo, empresa)",
            "Outro (descreva nas observações)",
        ],
    )

    add_section_divider(doc)

    # ── BLOCO 9: FUTURO E EVOLUÇÃO ────────────────────────────────────────────
    add_heading(doc, "9. Visão de Futuro e Evolução do Projeto", level=1)

    add_question_block(
        doc,
        "9.1",
        "Quais funcionalidades futuras você já imagina para o programa? "
        "Marque tudo que te interessa.",
        options=[
            "Identificar múltiplas pessoas ao mesmo tempo",
            "Processar vídeos além de fotos",
            "Integração com Google Fotos ou iCloud",
            "Interface mobile (celular ou tablet)",
            "Criar um banco de rostos com histórico de buscas",
            "Exportar resultados para planilha Excel",
            "Agendamento automático (rodar em horários programados)",
            "Nenhuma — quero algo simples e focado",
        ],
    )
    add_question_block(
        doc,
        "9.2",
        "Você tem alguma expectativa de prazo para ter uma primeira versão funcionando?",
        options=[
            "Sem pressa — quero que seja feito com qualidade",
            "Algumas semanas",
            "Preciso de algo básico funcionando o mais rápido possível",
        ],
    )
    add_question_block(
        doc,
        "9.3",
        "Existe alguma ideia, requisito ou preocupação que você tem sobre o projeto "
        "e que ainda não foi coberta pelas perguntas acima? Descreva livremente:",
        answer_lines=5,
    )

    add_section_divider(doc)

    # ── RODAPÉ ────────────────────────────────────────────────────────────────
    footer_para = doc.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_footer = footer_para.add_run(
        "Documento gerado para o projeto FacePhoto  •  Levantamento de Requisitos  •  v1.0"
    )
    run_footer.font.size = Pt(9)
    run_footer.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    run_footer.italic = True

    doc.save(output_path)
    print(f"[OK] Documento gerado com sucesso: {output_path}")


if __name__ == "__main__":
    output_file = (
        r"E:\Backup Trabalho e Hobbys\Cursos e Trabalhos"
        r"\PROJETOS PESSOAIS E TESTES\FacePhoto"
        r"\FacePhoto_Questionario_Requisitos.docx"
    )
    build_document(output_file)
